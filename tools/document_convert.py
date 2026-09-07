"""Office / PDF documents → Markdown for ``read_file`` (AIS-294).

Why: after a Teams/mail download the model's natural next move is
``read_file(saved_path)``. For ``.docx``/``.xlsx``/``.pptx`` that used to hit
the binary guard ("use terminal"), and a ``.pdf`` streamed raw bytes into the
context (support case SUP-20260907-073421). This module turns those files
into Markdown so the existing read path (pagination, line numbers, dedup)
just works.

Two backends, tried in the order configured by ``documents.converter``
(``auto`` = Suite first, then local):

* ``suite`` — the AIMDS-Suite's Docling behind ``go-mcp-customer``:
  ``storage_ingest_upload({})`` (returns the upload URL) → HTTP upload with
  the user's own LiteLLM virtual key → ``storage_ingest_upload({upload_id})``
  → ``storage_get_document({id: "upload:<id>", format: "markdown"})``. Only
  used when :func:`hermes_cli.iamds_suite.docling_availability` says the
  Suite is reachable (``/uptime/health`` answers, ``docling`` and
  ``customer-storage`` are up) and the storage tools are registered in this
  process. Any negative answer is logged once and falls through to local.
* ``local`` — ``markitdown`` with per-type fallbacks (``python-docx``,
  ``openpyxl``, ``python-pptx``, ``pypdf`` / PyMuPDF); the libraries of the
  ``office`` extra.

Results are cached under ``~/.hermes/cache/document_convert/`` keyed by the
resolved source path and invalidated on mtime/size change, so a re-read of a
document costs no conversion and the Vault stays free of generated sidecars.
"""

from __future__ import annotations

import concurrent.futures
import hashlib
import json
import logging
import os
import re
import threading
import time
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable, Dict, List, Optional

logger = logging.getLogger(__name__)

DOCUMENT_EXTENSIONS = frozenset({".docx", ".xlsx", ".pptx", ".odt", ".ods", ".odp", ".pdf"})
#: Pre-2007 binary Office formats: neither Docling nor the local libraries
#: read them — reported explicitly instead of silently falling back.
LEGACY_OFFICE_EXTENSIONS = frozenset({".doc", ".xls", ".ppt"})
#: Formats only the Suite (LibreOffice-backed Docling) can read.
_SUITE_ONLY_EXTENSIONS = frozenset({".odt", ".ods", ".odp"})

BACKEND_SUITE = "suite-docling"
BACKEND_LOCAL_MARKITDOWN = "local-markitdown"
BACKEND_LOCAL_DOCX = "local-python-docx"
BACKEND_LOCAL_XLSX = "local-openpyxl"
BACKEND_LOCAL_PPTX = "local-python-pptx"
BACKEND_LOCAL_PYPDF = "local-pypdf"
BACKEND_LOCAL_PYMUPDF = "local-pymupdf"

CONVERTER_AUTO = "auto"
CONVERTER_SUITE = "suite"
CONVERTER_LOCAL = "local"
_CONVERTER_MODES = (CONVERTER_AUTO, CONVERTER_SUITE, CONVERTER_LOCAL)

SUITE_MCP_SERVER = "AIMDSSuiteMCP"
_SUITE_INGEST_TOOL = "storage_ingest_upload"
_SUITE_GET_DOCUMENT_TOOL = "storage_get_document"
#: After a 404/5xx from the upload endpoint the Suite path is skipped for
#: this long — a Suite without the upload route enabled must not cost a
#: round trip per read.
SUITE_COOLDOWN_SECONDS = 600
_SUITE_UPLOAD_TIMEOUT = 120.0
_SUITE_UPLOAD_CONNECT_TIMEOUT = 15.0
_DEFAULT_MAX_UPLOAD_BYTES = 20 * 1024 * 1024

_LOCAL_TIMEOUT_SECONDS = 60.0
_LOCAL_MAX_TABLE_ROWS = 5000

_FRONTMATTER_RE = re.compile(r"\A---[ \t]*\r?\n.*?\r?\n---[ \t]*\r?\n", re.DOTALL)

_OFFICE_EXTRA_HINT = 'install the office extra (pip install -e ".[office]")'


class DocumentConvertError(Exception):
    """Raised when no backend could produce Markdown; ``reasons`` lists each attempt."""

    def __init__(self, message: str, reasons: Optional[List[str]] = None, suite_state: str = ""):
        super().__init__(message)
        self.reasons = list(reasons or [])
        self.suite_state = suite_state


class _SuiteUnavailable(Exception):
    """Suite backend cannot run right now (state carries the gate result)."""

    def __init__(self, state: str, reason: str):
        super().__init__(f"{state}: {reason}")
        self.state = state
        self.reason = reason


@dataclass
class ConvertResult:
    markdown: str
    backend: str
    cache_path: str
    source: str
    cached: bool = False
    suite_state: str = ""
    warnings: List[str] = field(default_factory=list)


# --------------------------------------------------------------------------- classification

def _suffix(path: str | os.PathLike[str]) -> str:
    return Path(str(path)).suffix.lower()


def is_convertible_document(path: str | os.PathLike[str]) -> bool:
    """True for Office/PDF files ``read_file`` returns as Markdown."""
    return _suffix(path) in DOCUMENT_EXTENSIONS


def is_legacy_office_document(path: str | os.PathLike[str]) -> bool:
    return _suffix(path) in LEGACY_OFFICE_EXTENSIONS


def legacy_document_error(path: str) -> str:
    return (
        f"Cannot read '{path}' ({_suffix(path)}): legacy binary Office format. "
        "Ask for a .docx/.xlsx/.pptx or PDF export — neither the AIMDS-Suite "
        "Docling nor the local converters read this format."
    )


# --------------------------------------------------------------------------- config

def converter_mode(config: Optional[dict] = None) -> str:
    """``documents.converter`` from config (env ``HERMES_DOCUMENT_CONVERTER`` wins)."""
    override = os.environ.get("HERMES_DOCUMENT_CONVERTER", "").strip().lower()
    if override in _CONVERTER_MODES:
        return override
    cfg = config
    if cfg is None:
        try:
            from hermes_cli.config import load_config

            cfg = load_config() or {}
        except Exception:
            cfg = {}
    section = cfg.get("documents") if isinstance(cfg, dict) else None
    mode = str((section or {}).get("converter") or "").strip().lower() if isinstance(section, dict) else ""
    return mode if mode in _CONVERTER_MODES else CONVERTER_AUTO


# --------------------------------------------------------------------------- cache

def cache_dir() -> Path:
    override = os.environ.get("HERMES_DOCUMENT_CACHE_DIR", "").strip()
    if override:
        return Path(override).expanduser()
    from hermes_constants import get_hermes_home

    return get_hermes_home() / "cache" / "document_convert"


def _cache_key(resolved: Path) -> str:
    return hashlib.sha1(str(resolved).encode("utf-8")).hexdigest()


def _cache_lookup(resolved: Path, stat: os.stat_result) -> Optional[ConvertResult]:
    key = _cache_key(resolved)
    base = cache_dir()
    meta_path, md_path = base / f"{key}.json", base / f"{key}.md"
    try:
        meta = json.loads(meta_path.read_text(encoding="utf-8"))
        if meta.get("source") != str(resolved):
            return None
        if int(meta.get("mtime_ns", -1)) != stat.st_mtime_ns or int(meta.get("size", -1)) != stat.st_size:
            return None
        markdown = md_path.read_text(encoding="utf-8")
    except (OSError, ValueError, TypeError):
        return None
    return ConvertResult(
        markdown=markdown,
        backend=str(meta.get("backend") or ""),
        cache_path=str(md_path),
        source=str(resolved),
        cached=True,
        suite_state=str(meta.get("suite_state") or ""),
    )


def _cache_store(resolved: Path, stat: os.stat_result, markdown: str, backend: str, suite_state: str) -> str:
    key = _cache_key(resolved)
    base = cache_dir()
    base.mkdir(parents=True, exist_ok=True)
    md_path = base / f"{key}.md"
    tmp = base / f"{key}.md.tmp"
    tmp.write_text(markdown, encoding="utf-8")
    os.replace(tmp, md_path)
    meta = {
        "source": str(resolved),
        "mtime_ns": stat.st_mtime_ns,
        "size": stat.st_size,
        "backend": backend,
        "suite_state": suite_state,
        "converted_at": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()),
    }
    (base / f"{key}.json").write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")
    return str(md_path)


# --------------------------------------------------------------------------- suite backend

_suite_lock = threading.Lock()
_suite_cooldown_until = 0.0
_suite_cooldown_reason = ""
_logged_suite_states: set[str] = set()


def _find_suite_tool(suffix: str):
    """Registered AIMDSSuiteMCP tool whose name ends with *suffix* (gateway prefix agnostic)."""
    try:
        from tools.registry import registry
    except Exception:
        return None
    prefix = f"mcp_{SUITE_MCP_SERVER}_"
    for entry in registry._snapshot_entries():
        name = getattr(entry, "name", "")
        if name.startswith(prefix) and name.endswith(suffix):
            return entry
    return None


def suite_tools_present() -> bool:
    return _find_suite_tool(_SUITE_INGEST_TOOL) is not None and _find_suite_tool(_SUITE_GET_DOCUMENT_TOOL) is not None


def _set_suite_cooldown(reason: str) -> None:
    global _suite_cooldown_until, _suite_cooldown_reason
    with _suite_lock:
        _suite_cooldown_until = time.monotonic() + SUITE_COOLDOWN_SECONDS
        _suite_cooldown_reason = reason


def reset_suite_cooldown() -> None:
    global _suite_cooldown_until, _suite_cooldown_reason
    with _suite_lock:
        _suite_cooldown_until = 0.0
        _suite_cooldown_reason = ""
    _logged_suite_states.clear()


def suite_availability(config: Optional[dict] = None):
    """:class:`hermes_cli.iamds_suite.DoclingAvailability` incl. the local cooldown."""
    from hermes_cli import iamds_suite as suite

    with _suite_lock:
        remaining = _suite_cooldown_until - time.monotonic()
        reason = _suite_cooldown_reason
    if remaining > 0:
        return suite.DoclingAvailability(
            suite.DOCLING_ENDPOINT_UNAVAILABLE,
            f"{reason} (retry in {int(remaining)}s)",
            checked_at=time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()),
        )
    return suite.docling_availability(config=config, tools_present=suite_tools_present)


def _log_suite_state_once(state: str, reason: str) -> None:
    if state in _logged_suite_states:
        return
    _logged_suite_states.add(state)
    logger.info("[AIS-294] Suite document conversion not used: %s (%s) — converting locally", state, reason)


def _parse_tool_output(raw: Any) -> Any:
    """Unwrap the JSON string an MCP tool handler returns.

    Handlers return ``{"result": <text>}`` (or ``{"error": ...}``); the text
    itself is usually JSON. Returns the innermost parsed value.
    """
    value = raw
    if isinstance(value, (bytes, bytearray)):
        value = value.decode("utf-8", "replace")
    if isinstance(value, str):
        try:
            value = json.loads(value)
        except ValueError:
            return value
    if isinstance(value, dict):
        if value.get("error"):
            raise RuntimeError(str(value["error"]))
        if "result" in value and len(value) <= 2:
            inner = value["result"]
            if isinstance(inner, str):
                try:
                    inner = json.loads(inner)
                except ValueError:
                    return inner
            return inner
    return value


def _call_suite_tool(suffix: str, args: Dict[str, Any]) -> Any:
    entry = _find_suite_tool(suffix)
    if entry is None:
        raise _SuiteUnavailable("tools_missing", f"{suffix} not registered")
    handler = getattr(entry, "handler", None)
    if handler is None:
        raise _SuiteUnavailable("tools_missing", f"{suffix} has no handler")
    return _parse_tool_output(handler(dict(args)))


def _strip_frontmatter(text: str) -> str:
    return _FRONTMATTER_RE.sub("", text, count=1)


def _document_text(payload: Any) -> str:
    if isinstance(payload, str):
        return payload
    if isinstance(payload, dict):
        for key in ("markdown", "content", "text", "full_text", "body"):
            value = payload.get(key)
            if isinstance(value, str) and value.strip():
                return value
        # MCP content blocks: [{"type": "text", "text": ...}]
        blocks = payload.get("content")
        if isinstance(blocks, list):
            texts = [b.get("text") for b in blocks if isinstance(b, dict) and isinstance(b.get("text"), str)]
            if texts:
                return "\n".join(texts)
    if isinstance(payload, list):
        texts = [b.get("text") for b in payload if isinstance(b, dict) and isinstance(b.get("text"), str)]
        if texts:
            return "\n".join(texts)
    return ""


def _upload_file(upload_url: str, resolved: Path, api_key: str, provider: str) -> str:
    """POST the file, return ``upload_id``. Maps HTTP failures to gate states."""
    import httpx

    from hermes_cli.iamds_suite import mark_suite_auth_failure

    headers = {"Authorization": f"Bearer {api_key}", "User-Agent": "hermes-agent/document-convert"}
    timeout = httpx.Timeout(_SUITE_UPLOAD_TIMEOUT, connect=_SUITE_UPLOAD_CONNECT_TIMEOUT)
    try:
        with resolved.open("rb") as fh:
            response = httpx.post(
                upload_url,
                files={"file": (resolved.name, fh, "application/octet-stream")},
                headers=headers,
                timeout=timeout,
                follow_redirects=False,
            )
    except httpx.HTTPError as exc:
        _set_suite_cooldown(f"upload failed: {type(exc).__name__}")
        raise _SuiteUnavailable("endpoint_unavailable", f"upload request failed: {exc}") from exc

    if response.status_code in (401, 403):
        mark_suite_auth_failure(provider, response.status_code, "document upload rejected", source="document_convert")
        raise _SuiteUnavailable("needs_reauth", f"upload returned HTTP {response.status_code}")
    if response.status_code == 404 or response.status_code >= 500:
        _set_suite_cooldown(f"upload endpoint returned HTTP {response.status_code}")
        raise _SuiteUnavailable("endpoint_unavailable", f"upload endpoint returned HTTP {response.status_code}")
    if response.status_code >= 400:
        raise _SuiteUnavailable("upload_rejected", f"upload returned HTTP {response.status_code}: {response.text[:200]}")
    try:
        body = response.json()
    except ValueError as exc:
        raise _SuiteUnavailable("upload_rejected", "upload endpoint returned no JSON") from exc
    upload_id = str((body or {}).get("upload_id") or "").strip()
    if not upload_id:
        raise _SuiteUnavailable("upload_rejected", "upload endpoint returned no upload_id")
    return upload_id


def _convert_via_suite(resolved: Path, stat: os.stat_result, config: Optional[dict]) -> tuple[str, str]:
    """Return ``(markdown, suite_state)``; raises :class:`_SuiteUnavailable`."""
    from hermes_cli.iamds_suite import resolve_suite_endpoint

    gate = suite_availability(config)
    if not gate.available:
        raise _SuiteUnavailable(gate.state, gate.reason)

    prepare = _call_suite_tool(_SUITE_INGEST_TOOL, {})
    if not isinstance(prepare, dict):
        raise _SuiteUnavailable("protocol", "storage_ingest_upload({}) returned no object")
    upload_url = str(prepare.get("upload_url") or "").strip()
    if not upload_url:
        raise _SuiteUnavailable("protocol", "storage_ingest_upload({}) returned no upload_url")
    try:
        max_bytes = int(prepare.get("max_size_bytes") or _DEFAULT_MAX_UPLOAD_BYTES)
    except (TypeError, ValueError):
        max_bytes = _DEFAULT_MAX_UPLOAD_BYTES
    if stat.st_size > max_bytes:
        raise _SuiteUnavailable("too_large", f"{stat.st_size} bytes exceeds the Suite limit of {max_bytes}")

    # Same environment as the active model: the gate resolved ``gate.provider``
    # from ``model.provider`` — its key (``gate.key_env``) signs the upload.
    ep = resolve_suite_endpoint(gate.provider, config=config, allow_default=True)
    if not ep.api_key:
        raise _SuiteUnavailable("needs_reauth", f"key_missing ({ep.key_env})")
    logger.info("[AIS-294] converting %s via Suite Docling (%s, %s)", resolved.name, ep.provider_id, ep.base_url)
    upload_id = _upload_file(upload_url, resolved, ep.api_key, ep.provider_id)

    _call_suite_tool(_SUITE_INGEST_TOOL, {"upload_id": upload_id})
    document = _call_suite_tool(_SUITE_GET_DOCUMENT_TOOL, {"id": f"upload:{upload_id}", "format": "markdown"})
    text = _strip_frontmatter(_document_text(document)).strip()
    if not text:
        raise _SuiteUnavailable("empty", "Suite returned no text for the document")
    return text + "\n", gate.state


# --------------------------------------------------------------------------- local backend

def _has_module(name: str) -> bool:
    import importlib.util

    try:
        return importlib.util.find_spec(name) is not None
    except (ImportError, ValueError):
        return False


def _pipe_table(rows: List[List[str]]) -> str:
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    norm = [[(c or "").replace("|", "\\|").replace("\n", " ").strip() for c in r] + [""] * (width - len(r)) for r in rows]
    header, body = norm[0], norm[1:]
    lines = ["| " + " | ".join(header) + " |", "|" + "|".join([" --- "] * width) + "|"]
    lines += ["| " + " | ".join(r) + " |" for r in body]
    return "\n".join(lines)


def _via_markitdown(resolved: Path) -> str:
    from markitdown import MarkItDown

    result = MarkItDown().convert(str(resolved))
    return (getattr(result, "text_content", None) or getattr(result, "markdown", "") or "").strip()


def _via_python_docx(resolved: Path) -> str:
    import docx

    document = docx.Document(str(resolved))
    parts: List[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name if paragraph.style is not None else "") or ""
        if style.startswith("Heading"):
            level = "".join(ch for ch in style if ch.isdigit()) or "1"
            parts.append("#" * min(int(level), 6) + " " + text)
        elif style.startswith("List"):
            parts.append("- " + text)
        else:
            parts.append(text)
    for table in document.tables:
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        parts.append(_pipe_table(rows))
    return "\n\n".join(p for p in parts if p).strip()


def _via_openpyxl(resolved: Path) -> str:
    import openpyxl

    workbook = openpyxl.load_workbook(str(resolved), read_only=True, data_only=True)
    parts: List[str] = []
    budget = _LOCAL_MAX_TABLE_ROWS
    for sheet in workbook.worksheets:
        rows: List[List[str]] = []
        for row in sheet.iter_rows(values_only=True):
            if budget <= 0:
                parts.append(f"_… truncated after {_LOCAL_MAX_TABLE_ROWS} rows_")
                break
            if row is None or all(v is None for v in row):
                continue
            rows.append(["" if v is None else str(v) for v in row])
            budget -= 1
        parts.append(f"## {sheet.title}\n\n" + (_pipe_table(rows) if rows else "_(empty sheet)_"))
    return "\n\n".join(parts).strip()


def _via_python_pptx(resolved: Path) -> str:
    from pptx import Presentation

    presentation = Presentation(str(resolved))
    parts: List[str] = []
    for index, slide in enumerate(presentation.slides, start=1):
        lines = [f"## Slide {index}"]
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text_frame is not None:
                for paragraph in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in paragraph.runs).strip()
                    if text:
                        lines.append(("  " * min(paragraph.level, 4)) + "- " + text)
            if getattr(shape, "has_table", False) and shape.table is not None:
                rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                lines.append(_pipe_table(rows))
        notes = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""
        if notes:
            lines.append(f"> Notes: {notes}")
        parts.append("\n".join(lines))
    return "\n\n".join(parts).strip()


def _via_pypdf(resolved: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(resolved))
    pages = []
    for index, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(f"## Page {index}\n\n{text}")
    return "\n\n".join(pages).strip()


def _via_pymupdf(resolved: Path) -> str:
    import fitz  # PyMuPDF

    pages = []
    with fitz.open(str(resolved)) as document:
        for index, page in enumerate(document, start=1):
            text = page.get_text().strip()
            if text:
                pages.append(f"## Page {index}\n\n{text}")
    return "\n\n".join(pages).strip()


_LOCAL_CHAIN: Dict[str, List[tuple[str, str, Callable[[Path], str]]]] = {
    ".docx": [("markitdown", BACKEND_LOCAL_MARKITDOWN, _via_markitdown), ("docx", BACKEND_LOCAL_DOCX, _via_python_docx)],
    ".xlsx": [("markitdown", BACKEND_LOCAL_MARKITDOWN, _via_markitdown), ("openpyxl", BACKEND_LOCAL_XLSX, _via_openpyxl)],
    ".pptx": [("markitdown", BACKEND_LOCAL_MARKITDOWN, _via_markitdown), ("pptx", BACKEND_LOCAL_PPTX, _via_python_pptx)],
    ".pdf": [("pypdf", BACKEND_LOCAL_PYPDF, _via_pypdf), ("fitz", BACKEND_LOCAL_PYMUPDF, _via_pymupdf), ("markitdown", BACKEND_LOCAL_MARKITDOWN, _via_markitdown)],
}


def _convert_locally(resolved: Path) -> tuple[str, str, List[str]]:
    """Return ``(markdown, backend, warnings)``; raises on total failure."""
    ext = resolved.suffix.lower()
    if ext in _SUITE_ONLY_EXTENSIONS:
        raise RuntimeError(f"{ext} needs the AIMDS-Suite Docling (no local converter)")
    chain = _LOCAL_CHAIN.get(ext)
    if not chain:
        raise RuntimeError(f"no local converter for {ext}")
    warnings: List[str] = []
    missing: List[str] = []
    for module, backend, fn in chain:
        if not _has_module(module):
            missing.append(module)
            continue
        try:
            with concurrent.futures.ThreadPoolExecutor(max_workers=1) as pool:
                text = pool.submit(fn, resolved).result(timeout=_LOCAL_TIMEOUT_SECONDS)
        except concurrent.futures.TimeoutError:
            warnings.append(f"{backend}: timed out after {int(_LOCAL_TIMEOUT_SECONDS)}s")
            continue
        except Exception as exc:  # noqa: BLE001 — try the next converter
            warnings.append(f"{backend}: {type(exc).__name__}: {str(exc)[:160]}")
            continue
        if text.strip():
            return text.rstrip() + "\n", backend, warnings
        warnings.append(f"{backend}: no text extracted")
    if missing and len(missing) == len(chain):
        raise RuntimeError(f"no converter library installed for {ext} — {_OFFICE_EXTRA_HINT}")
    detail = "; ".join(warnings) if warnings else "no converter succeeded"
    if ext == ".pdf" and all("no text" in w for w in warnings):
        detail += " (scanned PDF? OCR needs the AIMDS-Suite Docling)"
    raise RuntimeError(detail)


# --------------------------------------------------------------------------- entry point

def convert_document(path: str | os.PathLike[str], *, config: Optional[dict] = None, use_cache: bool = True) -> ConvertResult:
    """Markdown for an Office/PDF file, from cache, the Suite, or local libraries."""
    resolved = Path(str(path)).expanduser().resolve()
    if is_legacy_office_document(resolved):
        raise DocumentConvertError(legacy_document_error(str(path)), ["legacy format"])
    if not is_convertible_document(resolved):
        raise DocumentConvertError(f"'{path}' is not a convertible document ({resolved.suffix or 'no extension'})")
    try:
        stat = resolved.stat()
    except OSError as exc:
        raise DocumentConvertError(f"Cannot read '{path}': {exc.strerror or exc}") from exc

    if use_cache:
        hit = _cache_lookup(resolved, stat)
        if hit is not None:
            return hit

    mode = converter_mode(config)
    order = {CONVERTER_AUTO: (CONVERTER_SUITE, CONVERTER_LOCAL), CONVERTER_SUITE: (CONVERTER_SUITE,), CONVERTER_LOCAL: (CONVERTER_LOCAL,)}[mode]

    reasons: List[str] = []
    suite_state = ""
    warnings: List[str] = []
    for backend in order:
        if backend == CONVERTER_SUITE:
            try:
                markdown, suite_state = _convert_via_suite(resolved, stat, config)
            except _SuiteUnavailable as exc:
                suite_state = exc.state
                reasons.append(f"suite: {exc.state} — {exc.reason}")
                _log_suite_state_once(exc.state, exc.reason)
                continue
            except Exception as exc:  # noqa: BLE001 — never let the Suite path break a read
                suite_state = "error"
                reasons.append(f"suite: {type(exc).__name__}: {str(exc)[:200]}")
                logger.warning("[AIS-294] Suite document conversion failed for %s: %s", resolved.name, exc)
                continue
            cache_path = _cache_store(resolved, stat, markdown, BACKEND_SUITE, suite_state)
            return ConvertResult(markdown, BACKEND_SUITE, cache_path, str(resolved), suite_state=suite_state)
        else:
            try:
                markdown, used, warnings = _convert_locally(resolved)
            except Exception as exc:  # noqa: BLE001
                reasons.append(f"local: {exc}")
                continue
            cache_path = _cache_store(resolved, stat, markdown, used, suite_state)
            return ConvertResult(markdown, used, cache_path, str(resolved), suite_state=suite_state, warnings=warnings)

    raise DocumentConvertError(
        f"Could not convert '{path}' to text: " + "; ".join(reasons), reasons, suite_state=suite_state
    )
